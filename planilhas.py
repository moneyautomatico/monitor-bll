"""Planilhas: exportar (CSV que abre direto no Excel, e .xlsx quando a biblioteca existir) e importar (CSV ou .xlsx)."""
import csv
import io
import re
import unicodedata


def _norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def tem_xlsx():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def csv_bytes(cab, linhas):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(cab)
    for ln in linhas:
        w.writerow(["" if v is None else (f"{v:.2f}".replace(".", ",") if isinstance(v, float) else v) for v in ln])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM: o Excel brasileiro abre com acentos certos


def xlsx_bytes(cab, linhas, titulo="Planilha"):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo[:30]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    for ln in linhas:
        ws.append(list(ln))
    for col in ws.columns:
        larg = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, larg + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ler(nome, dados):
    """Lê CSV (; ou ,) ou .xlsx -> lista de dicts com o cabeçalho normalizado (sem acento, minúsculo)."""
    linhas = []
    if nome.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        ws = wb.active
        linhas = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        txt = None
        for cod in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                txt = dados.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        sep = ";" if (txt or "").split("\n", 1)[0].count(";") >= (txt or "").split("\n", 1)[0].count(",") else ","
        linhas = list(csv.reader(io.StringIO(txt or ""), delimiter=sep))
    linhas = [ln for ln in linhas if any(str(c).strip() for c in ln)]
    if not linhas:
        return []
    cab = [_norm(c) for c in linhas[0]]
    return [{cab[i]: str(v).strip() for i, v in enumerate(ln) if i < len(cab) and cab[i]} for ln in linhas[1:]]


def pegar(linha, *nomes):
    """Primeiro valor preenchido entre colunas com nomes parecidos (ex.: 'qtd', 'quantidade')."""
    for n in nomes:
        n = _norm(n)
        for k, v in linha.items():
            if (k == n or k.startswith(n)) and v:
                return v
    return ""
